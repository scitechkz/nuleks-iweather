"""
Extracts layout + formulas + styles of 'BOM Extract' and 'Weather Analysis'
from the Excel workbook into JSON used by the React app.
Run:  python scripts/extract_layout.py path/to/Inclement_Weather_Workbook.xlsx
"""
import sys, json, colorsys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

THEME = ['FFFFFF','000000','E7E6E6','44546A','5B9BD5','ED7D31','A5A5A5','FFC000','4472C4','70AD47','0563C1','954F72']
INDEXED = {8:'000000',9:'FFFFFF',10:'FF0000',11:'00FF00',12:'0000FF',13:'FFFF00',64:'000000'}

def tint(hexrgb, t):
    r,g,b=[int(hexrgb[i:i+2],16)/255 for i in (0,2,4)]
    h,l,s=colorsys.rgb_to_hls(r,g,b)
    l = l*(1+t) if t<0 else l*(1-t)+t
    r,g,b=colorsys.hls_to_rgb(h,l,s)
    return '%02X%02X%02X'%(round(r*255),round(g*255),round(b*255))

def color(c):
    if c is None: return None
    try:
        if c.type=='rgb' and c.rgb and isinstance(c.rgb,str):
            return '#'+c.rgb[-6:]
        if c.type=='theme':
            base=THEME[c.theme] if c.theme < len(THEME) else '000000'
            return '#'+(tint(base,c.tint) if c.tint else base)
        if c.type=='indexed':
            return '#'+INDEXED.get(c.indexed,'000000')
    except Exception:
        return None
    return None

def side(s):
    if s is None or s.style is None: return None
    return {'s':s.style,'c':color(s.color) or '#000000'}

def style_of(cell):
    f=cell.font; fl=cell.fill; b=cell.border; a=cell.alignment
    st={}
    fo={}
    if f.name: fo['n']=f.name
    if f.sz: fo['z']=f.sz
    if f.b: fo['b']=1
    if f.i: fo['i']=1
    if f.u: fo['u']=1
    col=color(f.color) if f.color is not None else None
    if col and col!='#000000': fo['c']=col
    if fo: st['f']=fo
    if fl is not None and fl.fill_type=='solid':
        bg=color(fl.fgColor)
        if bg: st['bg']=bg
    bd={k:side(getattr(b,k)) for k in ('left','right','top','bottom')}
    bd={k:v for k,v in bd.items() if v}
    if bd: st['b']=bd
    al={}
    if a.horizontal: al['h']=a.horizontal
    if a.vertical: al['v']=a.vertical
    if a.wrap_text: al['w']=1
    if a.text_rotation: al['r']=a.text_rotation
    if a.indent: al['i']=a.indent
    if al: st['a']=al
    if cell.number_format and cell.number_format!='General': st['nf']=cell.number_format
    return st

def extract(ws, name):
    styles=[]; sidx={}
    def sid(st):
        k=json.dumps(st,sort_keys=True)
        if k not in sidx:
            sidx[k]=len(styles); styles.append(st)
        return sidx[k]
    merges=[str(m) for m in ws.merged_cells.ranges]
    merged_covered=set()
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row,m.max_row+1):
            for c in range(m.min_col,m.max_col+1):
                if (r,c)!=(m.min_row,m.min_col): merged_covered.add((r,c))
    cells=[]
    for row in ws.iter_rows(min_row=1,max_row=ws.max_row,max_col=ws.max_column):
        for cell in row:
            st=style_of(cell)
            has=cell.value is not None
            if not has and not st: continue
            if (cell.row,cell.column) in merged_covered and not has:
                # still keep borders for merged areas
                if not st.get('b') and not st.get('bg'): continue
            v=cell.value
            d={'r':cell.row,'c':cell.column}
            if v is not None:
                if isinstance(v,str) and v.startswith('='): d['f']=v
                elif hasattr(v,'isoformat'):
                    d['v']=v.isoformat(); d['dt']=1
                else: d['v']=v
            d['s']=sid(st)
            cells.append(d)
    cols={}
    default_w = ws.sheet_format.defaultColWidth or 8.43
    for i in range(1,ws.max_column+1):
        L=get_column_letter(i)
        w=ws.column_dimensions[L].width if L in ws.column_dimensions and ws.column_dimensions[L].width else None
        cols[i]=w or default_w
    # openpyxl groups columns (min/max); expand ranges
    for L,cd in ws.column_dimensions.items():
        if cd.min and cd.max and cd.width:
            for i in range(cd.min,min(cd.max,ws.max_column)+1): cols[i]=cd.width
    rows={}
    default_h = ws.sheet_format.defaultRowHeight or 15
    for r in range(1,ws.max_row+1):
        rd=ws.row_dimensions.get(r)
        rows[r]=(rd.height if rd is not None and rd.height else default_h)
    return {'name':name,'maxRow':ws.max_row,'maxCol':ws.max_column,'merges':merges,
            'cols':cols,'rows':rows,'defaultRowHeight':default_h,'styles':styles,'cells':cells,
            'showGrid': ws.sheet_view.showGridLines is not False,
            'freezeRows': (ws.freeze_panes and int(''.join(ch for ch in ws.freeze_panes if ch.isdigit()))-1) or 0,
            'freezeCols': 0 if not ws.freeze_panes else (sum((ord(ch)-64)*26**i for i,ch in enumerate(reversed([c for c in ws.freeze_panes if c.isalpha()])))-1),
            'zoom': (ws.sheet_view.zoomScale or 100)}

wb=load_workbook(sys.argv[1])
out={}
for n in ('BOM Extract','Weather Analysis'):
    out[n]=extract(wb[n],n)
out['definedNames']={'HoursPerDayRatio':"'Weather Analysis'!$Q$24"}
json.dump(out,open(sys.argv[2] if len(sys.argv)>2 else 'src/data/layout.json','w'),separators=(',',':'))
print({k:(len(v['cells']),len(v['styles'])) for k,v in out.items() if isinstance(v,dict) and 'cells' in v})
for n in ('BOM Extract','Weather Analysis'):
    print(n,'showGrid',out[n]['showGrid'])
