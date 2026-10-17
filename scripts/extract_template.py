"""Extract the Excel template into a JSON model consumed by the React app.

Usage: python3 -I scripts/extract_template.py <template.xlsx> <out.json>
"""
import colorsys
import datetime
import json
import re
import sys
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

SRC, OUT = sys.argv[1], sys.argv[2]
SHEETS = ['BOM csv data', 'Climate', 'Weather Calculation', 'Critical Site Works Summary',
          'Rain+Other Factors', 'Project Site Calendar']

wb = openpyxl.load_workbook(SRC)
wbv = openpyxl.load_workbook(SRC, data_only=True)
zf = zipfile.ZipFile(SRC)

# ---------- theme colours ----------
theme_xml = zf.read('xl/theme/theme1.xml').decode('utf8')
scheme = {}
for name in ['dk1', 'lt1', 'dk2', 'lt2', 'accent1', 'accent2', 'accent3', 'accent4', 'accent5',
             'accent6', 'hlink', 'folHlink']:
    m = re.search(r'<a:%s>(.*?)</a:%s>' % (name, name), theme_xml, re.S)
    c = re.search(r'(?:lastClr|srgbClr val)="([0-9A-Fa-f]{6})"', m.group(1))
    scheme[name] = c.group(1).upper()
THEME = [scheme['lt1'], scheme['dk1'], scheme['lt2'], scheme['dk2'], scheme['accent1'], scheme['accent2'],
         scheme['accent3'], scheme['accent4'], scheme['accent5'], scheme['accent6'], scheme['hlink'],
         scheme['folHlink']]
from openpyxl.styles.colors import COLOR_INDEX


def apply_tint(hexrgb, tint):
    if not tint:
        return hexrgb
    r, g, b = (int(hexrgb[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tint) if tint < 0 else l * (1 - tint) + tint
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return '%02X%02X%02X' % (round(r * 255), round(g * 255), round(b * 255))


def color(c):
    if c is None:
        return None
    try:
        if c.type == 'rgb' and isinstance(c.rgb, str):
            v = c.rgb[-6:].upper()
            if c.rgb.upper().startswith('00') and len(c.rgb) == 8 and v == '000000':
                return v
            return apply_tint(v, c.tint)
        if c.type == 'theme':
            return apply_tint(THEME[c.theme], c.tint)
        if c.type == 'indexed':
            if c.indexed in (64, 65):
                return None  # system foreground/background
            return apply_tint(COLOR_INDEX[c.indexed][-6:].upper(), c.tint)
    except Exception:
        return None
    return None


def to_serial(v):
    if isinstance(v, datetime.datetime):
        d = v - datetime.datetime(1899, 12, 30)
        return d.days + d.seconds / 86400
    if isinstance(v, datetime.date):
        return (v - datetime.date(1899, 12, 30)).days
    return v


style_index = {}
styles = []


def style_of(c):
    f = c.font
    font = {'name': f.name, 'sz': f.sz, 'b': bool(f.b), 'i': bool(f.i), 'u': f.u, 'strike': bool(f.strike),
            'color': color(f.color)}
    fill = None
    if c.fill is not None and c.fill.fill_type == 'solid':
        fill = color(c.fill.fgColor) or color(c.fill.start_color)
    elif c.fill is not None and c.fill.fill_type not in (None, 'none'):
        fill = color(c.fill.fgColor)
    border = {}
    for side in ('left', 'right', 'top', 'bottom'):
        s = getattr(c.border, side)
        if s is not None and s.style:
            border[side[0]] = {'style': s.style, 'color': color(s.color) or '000000'}
    a = c.alignment
    align = {'h': a.horizontal, 'v': a.vertical, 'wrap': bool(a.wrap_text), 'indent': a.indent or 0,
             'rot': a.text_rotation or 0, 'shrink': bool(a.shrink_to_fit)}
    st = {'font': font, 'fill': fill, 'border': border, 'align': align, 'numFmt': c.number_format}
    key = json.dumps(st, sort_keys=True)
    if key not in style_index:
        style_index[key] = len(styles)
        styles.append(st)
    return style_index[key]


def dxf_of(dxf):
    if dxf is None:
        return None
    out = {}
    if dxf.fill is not None:
        out['fill'] = color(dxf.fill.bgColor) or color(dxf.fill.fgColor)
    if dxf.font is not None:
        out['color'] = color(dxf.font.color)
        if dxf.font.b:
            out['b'] = True
    return out


def drawing_for(sheet_index):
    rels = 'xl/worksheets/_rels/sheet%d.xml.rels' % sheet_index
    try:
        r = zf.read(rels).decode('utf8')
    except KeyError:
        return []
    m = re.search(r'Target="\.\./drawings/(drawing\d+\.xml)"', r)
    if not m:
        return []
    s = zf.read('xl/drawings/' + m.group(1)).decode('utf8')
    shapes = []
    for a in re.findall(r'<xdr:twoCellAnchor[^>]*>.*?</xdr:twoCellAnchor>', s, re.S):
        if '<xdr:graphicFrame' in a or '<xdr:pic' in a:
            continue
        fr = re.search(r'<xdr:from><xdr:col>(\d+)</xdr:col><xdr:colOff>(-?\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(-?\d+)</xdr:rowOff></xdr:from>', a)
        to = re.search(r'<xdr:to><xdr:col>(\d+)</xdr:col><xdr:colOff>(-?\d+)</xdr:colOff><xdr:row>(\d+)</xdr:row><xdr:rowOff>(-?\d+)</xdr:rowOff></xdr:to>', a)
        if not fr or not to:
            continue
        geom = re.search(r'prst="(\w+)"', a).group(1)
        link = re.search(r'textlink="([^"]*)"', a)
        txt = ''.join(re.findall(r'<a:t>([^<]*)</a:t>', a))
        line = re.search(r'<a:ln[^>]*>.*?<a:srgbClr val="(\w+)"', a, re.S)
        shapes.append({
            'name': re.search(r'name="([^"]*)"', a).group(1), 'geom': geom,
            'from': [int(x) for x in fr.groups()], 'to': [int(x) for x in to.groups()],
            'link': link.group(1).replace('$', '') if link and link.group(1) else None,
            'text': txt, 'flipH': 'flipH="1"' in a, 'flipV': 'flipV="1"' in a,
            'line': line.group(1) if line else None,
        })
    return shapes


out = {'theme': THEME, 'sheets': {}, 'styles': styles,
       'names': {n: v.attr_text for n, v in wb.defined_names.items()}}
for name in SHEETS:
    ws, wv = wb[name], wbv[name]
    idx = wb.sheetnames.index(name) + 1
    cells = {}
    max_r = max_c = 1
    from openpyxl.cell.cell import Cell
    blank_style = style_of(Cell(ws))
    for (r, c), cell in list(ws._cells.items()):
        v = cell.value
        s = style_of(cell)
        if v is None and s == blank_style:
            continue
        e = {'s': s}
        if isinstance(v, str) and v.startswith('='):
            e['f'] = v
            cv = wv.cell(row=r, column=c).value
            if cv is not None:
                e['cv'] = to_serial(cv)
        elif v is not None:
            e['v'] = to_serial(v)
        if cell.comment:
            e['note'] = cell.comment.text
        cells[cell.coordinate] = e
        if v is not None:
            max_r, max_c = max(max_r, r), max(max_c, c)
    cols = {}
    for key, d in ws.column_dimensions.items():
        lo, hi = d.min or column_index_from_string(key), d.max or column_index_from_string(key)
        for ci in range(lo, min(hi, 300) + 1):
            cols[ci] = {'w': d.width if d.customWidth or d.width else None, 'hidden': bool(d.hidden),
                        'outline': d.outlineLevel or 0}
    rows = {}
    for r, d in ws.row_dimensions.items():
        rows[r] = {'h': d.height, 'hidden': bool(d.hidden), 'outline': d.outlineLevel or 0}
    cf = []
    for rng, rules in ws.conditional_formatting._cf_rules.items():
        for rule in rules:
            e = {'sqref': str(rng.sqref), 'type': rule.type, 'priority': rule.priority,
                 'stop': bool(rule.stopIfTrue), 'operator': rule.operator, 'formula': list(rule.formula or []),
                 'dxf': dxf_of(rule.dxf)}
            if rule.colorScale is not None:
                e['cfvo'] = [{'type': v.type, 'val': v.val} for v in rule.colorScale.cfvo]
                e['colors'] = [color(c) for c in rule.colorScale.color]
            cf.append(e)
    sv = ws.sheet_view
    # dimension: include styled cells too
    for coord in cells:
        m = re.match(r'([A-Z]+)(\d+)', coord)
        max_c = max(max_c, column_index_from_string(m.group(1)))
        max_r = max(max_r, int(m.group(2)))
    out['sheets'][name] = {
        'cells': cells, 'cols': cols, 'rows': rows, 'merges': [str(m) for m in ws.merged_cells.ranges],
        'cf': cf, 'maxRow': max_r, 'maxCol': max_c,
        'defaultRowHeight': ws.sheet_format.defaultRowHeight or 15,
        'defaultColWidth': ws.sheet_format.defaultColWidth or 8.43,
        'showZeros': sv.showZeros is not False, 'showGrid': sv.showGridLines is not False,
        'zoom': sv.zoomScale or 100, 'shapes': drawing_for(idx),
    }

json.dump(out, open(OUT, 'w'), separators=(',', ':'))
print('styles', len(styles), {n: len(s['cells']) for n, s in out['sheets'].items()},
      {n: (s['maxRow'], s['maxCol']) for n, s in out['sheets'].items()})
